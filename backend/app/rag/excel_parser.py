from pathlib import Path
from typing import Any, Dict, List, Optional
from datetime import datetime
from zoneinfo import ZoneInfo
import json
import re

from openpyxl import load_workbook


BACKEND_ROOT = Path(__file__).resolve().parents[2]
SOURCE_DIR = BACKEND_ROOT / "knowledge_base" / "source_documents"
PARSED_DIR = BACKEND_ROOT / "knowledge_base" / "parsed"
DEFAULT_OUTPUT = PARSED_DIR / "parsed_rca_v4_excel.json"

MALAYSIA_TZ = ZoneInfo("Asia/Kuala_Lumpur")


def now_malaysia_iso() -> str:
    return datetime.now(MALAYSIA_TZ).replace(microsecond=0).isoformat()


def clean_text(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, str):
        value = value.replace("\r\n", "\n").replace("\r", "\n")
        value = re.sub(r"[ \t]+", " ", value)
        value = re.sub(r"\n{3,}", "\n\n", value)
        return value.strip()

    return value


def normalise_header(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).strip()
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text


def normalise_key(value: Any) -> str:
    text = normalise_header(value)
    text = text.replace("→", "to")
    text = text.replace("/", "_")
    text = re.sub(r"[^A-Za-z0-9]+", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.strip("_").lower()


def row_is_empty(row) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)


def find_header_row(ws, required_headers: List[str], max_scan_rows: int = 15) -> Optional[int]:
    required = {h.lower() for h in required_headers}

    for row_idx in range(1, min(ws.max_row, max_scan_rows) + 1):
        row_values = [
            normalise_header(ws.cell(row=row_idx, column=col_idx).value).lower()
            for col_idx in range(1, ws.max_column + 1)
        ]

        if required.issubset(set(row_values)):
            return row_idx

    return None


def sheet_to_records(ws, required_headers: List[str]) -> List[Dict[str, Any]]:
    header_row = find_header_row(ws, required_headers)

    if header_row is None:
        return []

    raw_headers = [
        normalise_header(ws.cell(row=header_row, column=col_idx).value)
        for col_idx in range(1, ws.max_column + 1)
    ]

    keys = [normalise_key(h) for h in raw_headers]

    records = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        row_values = [
            clean_text(ws.cell(row=row_idx, column=col_idx).value)
            for col_idx in range(1, ws.max_column + 1)
        ]

        if row_is_empty(row_values):
            continue

        record = {}
        for raw_header, key, value in zip(raw_headers, keys, row_values):
            if not key:
                continue
            record[key] = value
            record[f"{key}__header"] = raw_header

        record["_excel_row_number"] = row_idx
        records.append(record)

    return records


def read_key_value_sheet(ws, start_row: int = 1) -> List[Dict[str, Any]]:
    records = []

    for row_idx in range(start_row, ws.max_row + 1):
        left = clean_text(ws.cell(row=row_idx, column=1).value)
        right = clean_text(ws.cell(row=row_idx, column=2).value)

        if left is None and right is None:
            continue

        records.append({
            "key": left,
            "value": right,
            "_excel_row_number": row_idx,
        })

    return records


def parse_readme(ws) -> Dict[str, Any]:
    records = read_key_value_sheet(ws)

    title = records[0]["key"] if records else None
    description = records[1]["key"] if len(records) > 1 else None

    sheet_purpose = []
    header_row = find_header_row(ws, ["Sheet", "Purpose"])
    if header_row:
        for row_idx in range(header_row + 1, ws.max_row + 1):
            sheet = clean_text(ws.cell(row=row_idx, column=1).value)
            purpose = clean_text(ws.cell(row=row_idx, column=2).value)
            if sheet:
                sheet_purpose.append({
                    "sheet": sheet,
                    "purpose": purpose,
                    "_excel_row_number": row_idx,
                })

    return {
        "title": title,
        "description": description,
        "sheet_purpose": sheet_purpose,
    }


def parse_scoring_pipeline(ws) -> List[Dict[str, Any]]:
    """
    This sheet is explanatory. It is parsed as sequential text blocks.
    """

    blocks = []

    current_title = None
    current_lines = []

    for row_idx in range(1, ws.max_row + 1):
        value = clean_text(ws.cell(row=row_idx, column=1).value)
        if not value:
            continue

        is_title = (
            str(value).isupper()
            or str(value).startswith("STEP ")
            or "ERROR" in str(value).upper()
            or "SCORE" in str(value).upper()
        )

        if is_title and current_title is not None and current_lines:
            blocks.append({
                "title": current_title,
                "text": "\n\n".join(current_lines).strip(),
            })
            current_lines = []

        if is_title:
            current_title = str(value).split("\n")[0].strip()
            current_lines.append(str(value))
        else:
            if current_title is None:
                current_title = "Scoring Pipeline"
            current_lines.append(str(value))

    if current_title and current_lines:
        blocks.append({
            "title": current_title,
            "text": "\n\n".join(current_lines).strip(),
        })

    return blocks


def parse_score_model(ws) -> List[Dict[str, Any]]:
    return read_key_value_sheet(ws)


def find_source_excel(explicit_path: Optional[str | Path] = None) -> Path:
    if explicit_path:
        path = Path(explicit_path)
        if not path.exists():
            raise FileNotFoundError(f"Excel file not found: {path}")
        return path

    candidates = sorted(SOURCE_DIR.glob("*.xlsx"))

    if not candidates:
        raise FileNotFoundError(
            f"No .xlsx file found in {SOURCE_DIR}. "
            "Place Novaflow_RCA_V4_MAE_RMSE.xlsx there first."
        )

    # Prefer V4 file if multiple workbooks exist.
    for candidate in candidates:
        if "V4" in candidate.name.upper():
            return candidate

    return candidates[0]


def parse_rca_v4_excel(excel_path: Optional[str | Path] = None) -> Dict[str, Any]:
    path = find_source_excel(excel_path)

    wb = load_workbook(path, data_only=True)

    required_sheets = [
        "README",
        "1_SCORING_PIPELINE",
        "2_STAGE_DEFINITION",
        "3_RULES_MASTER",
        "4_THRESHOLD_LIBRARY",
        "5_ATTRIBUTION_ENGINE",
        "6_SCORE_MODEL",
        "7_RECOMMENDATION_MATRIX",
        "8_EXHAUST_SAFETY_CHECK",
    ]

    missing = [s for s in required_sheets if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"Missing required V4 sheet(s): {missing}")

    parsed = {
        "source_file": str(path),
        "source_file_name": path.name,
        "workbook_type": "Novaflow_RCA_V4_MAE_RMSE",
        "parsed_at": now_malaysia_iso(),
        "sheets": {
            "README": parse_readme(wb["README"]),
            "1_SCORING_PIPELINE": parse_scoring_pipeline(wb["1_SCORING_PIPELINE"]),
            "2_STAGE_DEFINITION": sheet_to_records(
                wb["2_STAGE_DEFINITION"],
                ["Stage_ID", "Stage_Name"],
            ),
            "3_RULES_MASTER": sheet_to_records(
                wb["3_RULES_MASTER"],
                ["Rule_ID", "Pri", "Stage", "Attribution", "Pattern"],
            ),
            "4_THRESHOLD_LIBRARY": sheet_to_records(
                wb["4_THRESHOLD_LIBRARY"],
                ["Rule_ID", "Metric_Name", "Unit"],
            ),
            "5_ATTRIBUTION_ENGINE": sheet_to_records(
                wb["5_ATTRIBUTION_ENGINE"],
                ["Step", "Priority"],
            ),
            "6_SCORE_MODEL": parse_score_model(wb["6_SCORE_MODEL"]),
            "7_RECOMMENDATION_MATRIX": sheet_to_records(
                wb["7_RECOMMENDATION_MATRIX"],
                ["Rec_Key", "Stage", "Root_Cause"],
            ),
            "8_EXHAUST_SAFETY_CHECK": sheet_to_records(
                wb["8_EXHAUST_SAFETY_CHECK"],
                ["Check_ID", "Check_Name"],
            ),
        },
    }

    parsed["summary"] = {
        "rules_count": len(parsed["sheets"]["3_RULES_MASTER"]),
        "threshold_count": len(parsed["sheets"]["4_THRESHOLD_LIBRARY"]),
        "recommendation_count": len(parsed["sheets"]["7_RECOMMENDATION_MATRIX"]),
        "stage_definition_count": len(parsed["sheets"]["2_STAGE_DEFINITION"]),
        "exhaust_safety_check_count": len(parsed["sheets"]["8_EXHAUST_SAFETY_CHECK"]),
    }

    return parsed


def save_parsed_excel(parsed: Dict[str, Any], output_path: str | Path = DEFAULT_OUTPUT) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    with output.open("w", encoding="utf-8") as f:
        json.dump(parsed, f, ensure_ascii=False, indent=2)

    return output


def parse_and_save_rca_v4_excel(
    excel_path: Optional[str | Path] = None,
    output_path: str | Path = DEFAULT_OUTPUT,
) -> Dict[str, Any]:
    parsed = parse_rca_v4_excel(excel_path)
    saved_path = save_parsed_excel(parsed, output_path)

    return {
        "source_file": parsed["source_file"],
        "output_path": str(saved_path),
        **parsed["summary"],
        "parsed_at": parsed["parsed_at"],
    }


if __name__ == "__main__":
    result = parse_and_save_rca_v4_excel()
    print(json.dumps(result, indent=2))
